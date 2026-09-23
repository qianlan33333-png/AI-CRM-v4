"""Excel preparation and observation service.

The component owns the upload/observation boundary only. It never approves
or sends messages and it never derives a customer or an audience segment from
an external source. The Go AI Assistant bridge supplies the final rows and
the approved, immutable segmentation snapshot.
"""

import argparse
import csv
import hashlib
import hmac
import io
import json
import os
import re
import sqlite3
import threading
import urllib.parse
import zipfile
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from zoneinfo import ZoneInfo


HEADERS = ["unionid", "话术", "小程序 path", "发送人 userid", "标题"]
SEGMENT_HEADER = "分层"
ALLOWED_SEGMENTS = frozenset("ABCD")
UNSEGMENTED = "unsegmented"
WINDOWS = (12, 24, 48)
MAX_BYTES = 8 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_FILES = 2000
MAX_ROWS = 5000
MAX_COVER_BYTES = 2 * 1024 * 1024
SHANGHAI = ZoneInfo("Asia/Shanghai")
REPORT_INSTANT = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d{1,9})?(?:Z|[+-]\d{2}:\d{2})$")
DELIVERY_LABELS = {
    "pending_submission": "待提交",
    "task_created_waiting_employee": "任务已创建，待员工执行",
    "provider_accepted": "任务已创建，待员工执行",
    "delivery_proven": "发送成功",
    "final_failed": "明确失败",
    "outcome_unknown": "结果待核实",
}
WINDOW_LABELS = {
    "observing": "观察中",
    "opened": "已打开",
    "not_opened": "未打开",
    "unavailable": "暂不可统计",
}


def stamp():
    return datetime.now(timezone.utc).isoformat()


def instant(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    if not isinstance(value, str):
        raise ValueError("timestamp must be text")
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(timezone.utc)


def digest(raw):
    return "sha256:" + hashlib.sha256(raw).hexdigest()


class Invalid(ValueError):
    """Input validation error, optionally with structured row errors."""

    def __init__(self, message, row_errors=None):
        super().__init__(message)
        self.row_errors = list(row_errors or [])


class Conflict(Exception):
    """A safe-to-report command conflict."""

    def __init__(self, message, code="conflict"):
        super().__init__(message)
        self.code = code


def _row_error(row, field, code):
    return {"row": row, "field": field, "code": code}


def _validate_segment(value, row_number=None):
    if value is None:
        return None
    if not isinstance(value, str):
        message = "分层必须是 A/B/C/D 或空"
        errors = [_row_error(row_number, SEGMENT_HEADER, "invalid_segment")] if row_number else []
        raise Invalid(message, errors)
    value = value.strip()
    if value == "":
        return None
    if value not in ALLOWED_SEGMENTS:
        message = "分层必须是 A/B/C/D 或空"
        errors = [_row_error(row_number, SEGMENT_HEADER, "invalid_segment")] if row_number else []
        raise Invalid(message, errors)
    return value


def _header_indexes(cells):
    """Return a name-to-column map while tolerating formatted blank tails."""
    values = [cell.value for cell in cells]
    while values and values[-1] is None:
        values.pop()
    if not values:
        raise Invalid("Excel 首行不能为空")
    for value in values:
        if not isinstance(value, str) or value == "":
            raise Invalid("Excel 首行只能包含已知列名")
    duplicates = {value for value in values if values.count(value) > 1}
    if duplicates:
        raise Invalid("Excel 首行存在重复列名：" + "、".join(sorted(duplicates)))
    unknown = [value for value in values if value not in set(HEADERS) | {SEGMENT_HEADER}]
    if unknown:
        raise Invalid("Excel 存在未知列：" + "、".join(unknown))
    missing = [name for name in HEADERS if name not in values]
    if missing:
        raise Invalid("Excel 缺少必需列：" + "、".join(missing))
    return {value: index for index, value in enumerate(values)}, SEGMENT_HEADER in values


def _cell_value(cell, field, row_number, allow_empty=True):
    value = cell.value
    if cell.data_type == "f":
        raise Invalid(
            f"第 {row_number} 行：{field} 不能使用公式",
            [_row_error(row_number, field, "formula_not_allowed")],
        )
    if value is None and allow_empty:
        return None
    if not isinstance(value, str):
        raise Invalid(
            f"第 {row_number} 行：{field} 必须是文本",
            [_row_error(row_number, field, "text_required")],
        )
    return value


def parse_excel_details(raw):
    """Parse an XLSX by column name and return rows plus segmentation metadata."""
    from openpyxl import load_workbook

    if not raw or len(raw) > MAX_BYTES:
        raise Invalid("文件为空或超过 8 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES or len(archive.infolist()) > MAX_ARCHIVE_FILES:
                raise Invalid("Excel 解压后过大")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Invalid:
        raise
    except Exception:
        raise Invalid("请上传有效的 .xlsx 文件") from None

    try:
        if not book.worksheets:
            raise Invalid("Excel 没有工作表")
        sheet = book.worksheets[0]
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise Invalid(f"每批最多 {MAX_ROWS} 行")
        rows = sheet.iter_rows()
        header_cells = next(rows, ())
        indexes, has_segment_column = _header_indexes(header_cells)
        max_index = max(indexes.values())
        result, seen, text_bytes = [], set(), 0

        for number, cells in enumerate(rows, 2):
            if not any(cell.value is not None for cell in cells):
                continue
            # A value after the last named column is data under an unnamed
            # column. Reject it instead of silently dropping user input.
            if len(cells) > max_index + 1 and any(cell.value is not None for cell in cells[max_index + 1 :]):
                raise Invalid(f"第 {number} 行：存在未命名列")

            values = {}
            for name in HEADERS:
                index = indexes[name]
                cell = cells[index] if index < len(cells) else None
                values[name] = None if cell is None else _cell_value(cell, name, number, allow_empty=name == "标题")
            if has_segment_column:
                index = indexes[SEGMENT_HEADER]
                cell = cells[index] if index < len(cells) else None
                values[SEGMENT_HEADER] = None if cell is None else _cell_value(cell, SEGMENT_HEADER, number)

            unionid = values["unionid"]
            text = values["话术"]
            path = values["小程序 path"]
            sender = values["发送人 userid"]
            title = values["标题"] or ""
            for name, value in (("unionid", unionid), ("话术", text), ("小程序 path", path), ("发送人 userid", sender)):
                if value is None or not value.strip():
                    raise Invalid(
                        f"第 {number} 行：{name} 不能为空",
                        [_row_error(number, name, "required")],
                    )
            unionid, path, sender = unionid.strip(), path.strip(), sender.strip()
            title = title.strip()
            segment = _validate_segment(values.get(SEGMENT_HEADER), number) if has_segment_column else None
            if any(re.search(r"[\x00-\x20]", value) for value in (unionid, path, sender)):
                raise Invalid(
                    f"第 {number} 行：存在空白或无效标识",
                    [_row_error(number, "unionid", "invalid_identifier")],
                )
            if len(unionid.encode()) > 256 or len(sender.encode()) > 256 or len(path.encode()) > 1024 or len(text.encode()) > 8000 or len(title.encode()) > 512:
                raise Invalid(f"第 {number} 行：内容超长", [_row_error(number, "row", "too_long")])
            if "://" in path or path.startswith("//") or not path.startswith("pages/"):
                raise Invalid(f"第 {number} 行：需要完整的小程序 pages/ 路径", [_row_error(number, "小程序 path", "invalid_path")])
            if unionid in seen:
                raise Invalid(
                    f"第 {number} 行：同一批次的 UnionID 重复，请先合并为一行",
                    [_row_error(number, "unionid", "duplicate_unionid")],
                )
            text_bytes += sum(len(value.encode()) for value in (unionid, text, path, sender, title, segment or ""))
            if text_bytes > MAX_BYTES:
                raise Invalid("单批次文本总量不能超过 8 MB")
            seen.add(unionid)
            result.append({
                "unionid": unionid,
                "text": text,
                "path": path,
                "sender_userid": sender,
                "title": title,
                "segment": segment,
            })
            if len(result) > MAX_ROWS:
                raise Invalid(f"每批最多 {MAX_ROWS} 行")
        if not result:
            raise Invalid("Excel 没有接收用户")
        return result, has_segment_column, any(row["segment"] for row in result)
    finally:
        book.close()


def parse_excel(raw):
    """Backward-compatible row-only parser."""
    return parse_excel_details(raw)[0]


def content_key(path):
    """Only known routes are measurable. Never treat a generic app visit as an open."""
    parsed = urllib.parse.urlsplit(path)
    query = urllib.parse.parse_qs(parsed.query)
    if parsed.path == "pages/article/article" and len(query.get("lesson_id", [])) == 1:
        return ("lesson", query["lesson_id"][0])
    if parsed.path in ("pages/case/case", "pages/case-detail/case-detail") and len(query.get("case_id", [])) == 1:
        return ("case", query["case_id"][0])
    return None


class Source:
    """Read-only open-event adapter; segmentation is never queried here."""

    def __init__(self, config):
        self.config = config
        self.local = threading.local()

    @contextmanager
    def session(self):
        self.local.batch = True
        self.local.connection = None
        try:
            yield
        finally:
            if self.local.connection is not None:
                self.local.connection.close()
            self.local.connection = None
            self.local.batch = False

    def query(self, name, params):
        import pymysql

        query = self.config.get(name)
        if not query or not query.lstrip().upper().startswith("SELECT "):
            raise RuntimeError("source_unconfigured")
        batch = getattr(self.local, "batch", False)
        db = getattr(self.local, "connection", None) if batch else None
        if db is None:
            db = pymysql.connect(
                **self.config["mysql"],
                cursorclass=pymysql.cursors.DictCursor,
                connect_timeout=5,
                read_timeout=20,
                write_timeout=5,
                autocommit=False,
            )
            if batch:
                self.local.connection = db
        try:
            with db.cursor() as cursor:
                cursor.execute("SET TRANSACTION READ ONLY")
                cursor.execute("START TRANSACTION")
                cursor.execute(query, params)
                # Full result for this recipient/content/window; no global LIMIT.
                return cursor.fetchall()
        finally:
            db.rollback()
            if not batch:
                db.close()

    def opens(self, unionid, path, start, end):
        key = content_key(path)
        if not key:
            return None
        try:
            # Empty event rows prove zero opens only when the collector confirms
            # complete coverage of this content and interval.
            coverage = self.query("coverage_sql", (key[0], key[1], start, end))
            if len(coverage) != 1 or coverage[0].get("complete") != 1:
                return None
            users = self.query("user_sql", (unionid,))
            if len(users) != 1:
                return None
            rows = self.query(key[0] + "_opens_sql", (users[0]["user_id"], key[1], start, end))
            return [instant(row["opened_at"]).isoformat() for row in rows]
        except Exception:
            return None


def _empty_stats():
    return {
        "sent": 0,
        "matured": 0,
        "observing": 0,
        "opened": 0,
        "unavailable": 0,
        "open_rate": None,
    }


def _copy_stats(stats):
    return dict(stats)


def _finalize_stats(stats):
    if stats["matured"] == 0 or stats["unavailable"]:
        stats["open_rate"] = None
    else:
        stats["open_rate"] = stats["opened"] / stats["matured"]
    return stats


def _safe_csv(value):
    text = "" if value is None else str(value)
    if text and text[0] in "=+-@":
        return "'" + text
    return text


def _report_time(value):
    """Format a Go RFC3339 delivery instant for a business-facing CSV.

    Observations retain their original UTC text and all observation arithmetic
    stays in UTC. Only this download boundary is localized. The bridge
    supplies Go time.Time JSON, so a naive historical string is not an
    established source contract and must not be guessed as local or UTC.
    """
    if value is None or value == "":
        return ""
    if not isinstance(value, str) or not REPORT_INSTANT.fullmatch(value):
        return "时间暂时无法显示"
    try:
        return instant(value).astimezone(SHANGHAI).strftime("%Y-%m-%d %H:%M:%S")
    except (OverflowError, ValueError):
        return "时间暂时无法显示"


def _report_label(value, labels, empty=""):
    if value is None or value == "":
        return empty
    return labels.get(value, "状态待核对") if isinstance(value, str) else "状态待核对"


def _normalize_rows(rows):
    normalized = []
    for row in rows:
        item = dict(row)
        segment = _validate_segment(item.get("segment"))
        item["segment"] = segment
        card = dict(item.get("card") or {})
        normalized_card = {
            "appid": card.get("appid", ""),
            "path": card.get("path", item.get("path", "")),
            "title": card.get("title", item.get("title", "")) or "",
            "cover_digest": card.get("cover_digest", "") or "",
        }
        item["card"] = normalized_card
        normalized.append(item)
    return normalized


class Service:
    def __init__(self, database, config, source=None):
        self.database, self.config = str(database), config
        self.source = source or Source(config.get("source", {}))
        with self.db() as db:
            db.executescript(
                """
                PRAGMA journal_mode=WAL;
                CREATE TABLE IF NOT EXISTS imports (
                  batch_key TEXT PRIMARY KEY,
                  file_digest TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  rows_json TEXT NOT NULL,
                  plan_id INTEGER UNIQUE,
                  request_key TEXT UNIQUE,
                  segment_column_present INTEGER NOT NULL DEFAULT 0,
                  has_segments INTEGER NOT NULL DEFAULT 0
                );
                CREATE INDEX IF NOT EXISTS imports_digest ON imports(file_digest,created_at);
                CREATE TABLE IF NOT EXISTS covers (
                  digest TEXT PRIMARY KEY,
                  body BLOB NOT NULL,
                  mime TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS snapshots (
                  snapshot_key TEXT PRIMARY KEY,
                  created_at TEXT NOT NULL,
                  groups_json TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS observations (
                  plan_id INTEGER PRIMARY KEY,
                  body TEXT NOT NULL,
                  updated_at TEXT NOT NULL
                );
                """
            )
            self._ensure_column(db, "imports", "segment_column_present", "INTEGER NOT NULL DEFAULT 0")
            self._ensure_column(db, "imports", "has_segments", "INTEGER NOT NULL DEFAULT 0")

    @staticmethod
    def _ensure_column(db, table, column, definition):
        columns = {row[1] for row in db.execute(f"PRAGMA table_info({table})")}
        if column not in columns:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")

    @contextmanager
    def db(self):
        db = sqlite3.connect(self.database, timeout=30)
        db.execute("PRAGMA foreign_keys=ON")
        db.row_factory = sqlite3.Row
        try:
            with db:
                yield db
        finally:
            db.close()

    def _imported(self, row, replayed):
        rows = self._rows_for_import(row)
        has_segments = bool(row["has_segments"] or any(item.get("segment") for item in rows))
        response = {
            "batch_key": row["batch_key"],
            "file_digest": row["file_digest"],
            "created_at": row["created_at"],
            "plan_id": row["plan_id"],
            "version": 1,
            "segment_source": "excel",
            "segment_column_present": bool(row["segment_column_present"]),
            "has_segments": has_segments,
            "rows": rows,
            "summary": {
                "total_rows": len(rows),
                "excluded_rows": 0,
                "empty_title_rows": sum(not item.get("title", "") for item in rows),
                "expected_tasks": len(rows),
            },
            "replayed": replayed,
        }
        return response

    def _rows_for_import(self, row):
        return _normalize_rows(json.loads(row["rows_json"]))

    def cover(self, raw):
        if len(raw) > MAX_COVER_BYTES:
            raise Invalid("封面超过 2 MB")
        if raw.startswith(b"\x89PNG\r\n\x1a\n"):
            mime = "image/png"
        elif raw.startswith(b"\xff\xd8\xff"):
            mime = "image/jpeg"
        else:
            raise Invalid("封面必须是 PNG 或 JPEG")
        key = digest(raw)
        with self.db() as db:
            db.execute("INSERT OR IGNORE INTO covers VALUES(?,?,?)", (key, raw, mime))
        return key

    def prepare_import(self, raw):
        rows, has_segment_column, has_segments = parse_excel_details(raw)
        for row in rows:
            row["card"] = {
                "appid": self.config.get("appid", ""),
                "path": row["path"],
                "title": row["title"],
                "cover_digest": "",
            }
        return {
            "file_digest": digest(raw),
            "segment_source": "excel",
            "segment_column_present": has_segment_column,
            "has_segments": has_segments,
            "rows": rows,
        }

    def import_file(self, raw, new_batch=False, request_key=""):
        if new_batch and not re.fullmatch(r"[A-Za-z0-9:_-]{8,200}", request_key or ""):
            raise Invalid("新批次需要稳定的请求编号")
        file_digest = digest(raw)
        with self.db() as db:
            prior = db.execute("SELECT * FROM imports WHERE request_key=?", (request_key or None,)).fetchone()
            if prior and prior["file_digest"] != file_digest:
                raise Conflict("同一请求编号不能更换文件", "idempotency_conflict")
            if not prior and not new_batch:
                prior = db.execute("SELECT * FROM imports WHERE file_digest=? ORDER BY created_at LIMIT 1", (file_digest,)).fetchone()
            if prior:
                response = self._imported(prior, True)
                return response

        prepared = self.prepare_import(raw)
        batch_key = hashlib.sha256((file_digest + (request_key if new_batch else "")).encode()).hexdigest()
        created_at = stamp()
        rows_json = json.dumps(prepared["rows"], ensure_ascii=False)
        with self.db() as db:
            existing = db.execute("SELECT * FROM imports WHERE batch_key=?", (batch_key,)).fetchone()
            if existing:
                response = self._imported(existing, True)
                return response
            db.execute(
                "INSERT INTO imports(batch_key,file_digest,created_at,rows_json,request_key,segment_column_present,has_segments) VALUES(?,?,?,?,?,?,?)",
                (batch_key, file_digest, created_at, rows_json, request_key or None, int(prepared["segment_column_present"]), int(prepared["has_segments"])),
            )
            result = db.execute("SELECT * FROM imports WHERE batch_key=?", (batch_key,)).fetchone()
            return self._imported(result, False)

    def _snapshot_payload(self, rows, segment_source=None, has_segments=None, version=None):
        groups = {}
        seen = set()
        explicit = any("segment" in row for row in rows)
        if segment_source is None:
            segment_source = "excel" if explicit else "unavailable"
        if segment_source not in ("excel", "unavailable"):
            raise Invalid("segment_source 无效")
        for index, row in enumerate(rows, 1):
            row_id = row.get("id")
            if row_id is None:
                raise Invalid(f"第 {index} 行缺少 id")
            key = str(row_id)
            if key in seen:
                raise Invalid(f"第 {index} 行重复 id")
            seen.add(key)
            segment = _validate_segment(row.get("segment"))
            if explicit:
                # Keep empty input rows too. Once any row has a real segment,
                # these rows become the frozen ``unsegmented`` bucket.
                groups[key] = segment
            elif segment:
                groups[key] = segment
        inferred = any(value for value in groups.values())
        if has_segments is None:
            has_segments = inferred
        if not isinstance(has_segments, bool):
            raise Invalid("has_segments 必须是布尔值")
        if has_segments and not inferred:
            raise Invalid("没有有效分层值")
        if not has_segments:
            groups = {}
        return {"version": version, "segment_source": segment_source, "has_segments": has_segments, "groups": groups}

    def snapshot(self, key, rows, segment_source=None, has_segments=None, version=None):
        payload = self._snapshot_payload(rows, segment_source, has_segments, version)
        serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True)
        with self.db() as db:
            prior = db.execute("SELECT created_at,groups_json FROM snapshots WHERE snapshot_key=?", (key,)).fetchone()
            if prior:
                if prior["groups_json"] != serialized:
                    raise Conflict("分层快照已冻结，不能替换", "snapshot_conflict")
                return {"snapshot_key": key, "created_at": prior["created_at"], **payload}
            created_at = stamp()
            db.execute("INSERT INTO snapshots VALUES(?,?,?)", (key, created_at, serialized))
        return {"snapshot_key": key, "created_at": created_at, **payload}

    @staticmethod
    def _read_snapshot(raw):
        try:
            payload = json.loads(raw)
        except Exception:
            payload = {}
        if isinstance(payload, dict) and "groups" in payload:
            groups = payload.get("groups") if isinstance(payload.get("groups"), dict) else {}
            return {"version": payload.get("version"), "segment_source": payload.get("segment_source", "unavailable"), "has_segments": bool(payload.get("has_segments")), "groups": groups if payload.get("has_segments") else {}}
        # Old snapshots cannot be represented as Excel segmentation, so report
        # them as overall-only instead of inventing an unknown group.
        return {"version": None, "segment_source": "unavailable", "has_segments": False, "groups": {}}

    def observe(self, plan_id, snapshot_key, rows, now=None):
        now = instant(now) if now else datetime.now(timezone.utc)
        with self.db() as db:
            snapshot = db.execute("SELECT groups_json FROM snapshots WHERE snapshot_key=?", (snapshot_key,)).fetchone()
        frozen = self._read_snapshot(snapshot["groups_json"]) if snapshot else {"version": None, "segment_source": "unavailable", "has_segments": False, "groups": {}}
        groups = frozen["groups"]
        has_segments = frozen["has_segments"]
        present_groups = set()
        details, counts = [], {}
        for row in rows:
            item = dict(row)
            state = item.get("state", item.get("delivery_state", ""))
            counts[state] = counts.get(state, 0) + 1
            segment = groups.get(str(item.get("id"))) if has_segments else None
            item["segment"] = segment
            item["windows"] = {}
            if has_segments:
                present_groups.add(segment or UNSEGMENTED)
            details.append(item)
        ordered_groups = [name for name in "ABCD" if name in present_groups]
        if has_segments and UNSEGMENTED in present_groups:
            ordered_groups.append(UNSEGMENTED)

        report_windows = {}
        overall_report = {}
        for hours in WINDOWS:
            overall = _empty_stats()
            group_stats = {name: _empty_stats() for name in ordered_groups}
            for item in details:
                state = item.get("state", item.get("delivery_state", ""))
                sent_at = item.get("sent_at")
                if state != "delivery_proven" or not sent_at:
                    continue
                try:
                    sent = instant(sent_at)
                except Exception:
                    continue
                end = sent + timedelta(hours=hours)
                overall["sent"] += 1
                target_group = item.get("segment") or UNSEGMENTED
                group = group_stats.get(target_group)
                if group is not None:
                    group["sent"] += 1
                if now < end:
                    overall["observing"] += 1
                    if group is not None:
                        group["observing"] += 1
                    item["windows"][str(hours)] = "observing"
                    continue
                overall["matured"] += 1
                if group is not None:
                    group["matured"] += 1
                events = self.source.opens(item.get("unionid", ""), item.get("path", ""), sent, end)
                if events is None:
                    overall["unavailable"] += 1
                    if group is not None:
                        group["unavailable"] += 1
                    item["windows"][str(hours)] = "unavailable"
                    continue
                try:
                    times = [instant(value) for value in events]
                except Exception:
                    times = None
                if times is None:
                    overall["unavailable"] += 1
                    if group is not None:
                        group["unavailable"] += 1
                    item["windows"][str(hours)] = "unavailable"
                    continue
                opened = any(sent <= value <= end for value in times)
                if opened:
                    overall["opened"] += 1
                    if group is not None:
                        group["opened"] += 1
                item["windows"][str(hours)] = "opened" if opened else "not_opened"
            _finalize_stats(overall)
            for stats in group_stats.values():
                _finalize_stats(stats)
            report_windows[str(hours)] = {"overall": overall, "groups": group_stats if has_segments else {}, "has_segments": has_segments}
            # Keep the top-level window map for clients that render a compact
            # overall summary, while the per-window object carries groups.
            overall_report[str(hours)] = _copy_stats(overall)
        result = {
            "plan_id": plan_id,
            "snapshot_key": snapshot_key,
            "updated_at": now.isoformat(),
            "collected_at": now.isoformat(),
            "observed_at": now.isoformat(),
            "segment_source": frozen["segment_source"],
            "has_segments": has_segments,
            "counts": counts,
            "overall": overall_report,
            "windows": report_windows,
            "rows": details,
        }
        with self.db() as db:
            db.execute(
                "INSERT INTO observations VALUES(?,?,?) ON CONFLICT(plan_id) DO UPDATE SET body=excluded.body,updated_at=excluded.updated_at",
                (plan_id, json.dumps(result, ensure_ascii=False), now.isoformat()),
            )
        return result

    def content_open(self, data):
        unionid = data.get("unionid", "")
        path = data.get("path", "")
        if not isinstance(unionid, str) or not isinstance(path, str) or not unionid or len(unionid.encode()) > 256 or len(path.encode()) > 1024:
            raise Invalid("内容打开查询参数无效")
        try:
            start, end = instant(data["start"]), instant(data["end"])
        except Exception as error:
            raise Invalid("内容打开查询时间无效") from error
        if end <= start or end - start > timedelta(hours=24, minutes=1):
            raise Invalid("内容打开查询窗口无效")
        if content_key(path) is None:
            return {"state": "unavailable", "reason": "unsupported_path"}
        events = self.source.opens(unionid, path, start, end)
        if events is None:
            return {"state": "unavailable", "reason": "coverage_incomplete"}
        try:
            opened = sorted(value for value in (instant(event) for event in events) if start <= value <= end)
        except Exception:
            return {"state": "unavailable", "reason": "invalid_source_event"}
        if opened:
            return {"state": "opened", "opened_at": opened[0].isoformat()}
        return {"state": "not_opened"}

    def report(self, plan_id):
        with self.db() as db:
            row = db.execute("SELECT body FROM observations WHERE plan_id=?", (plan_id,)).fetchone()
        if row:
            return json.loads(row["body"])
        return {
            "plan_id": plan_id,
            "pending": True,
            "segment_source": "unavailable",
            "has_segments": False,
            "observed_at": None,
            "overall": {str(hours): _empty_stats() for hours in WINDOWS},
            "windows": {
                str(hours): {"overall": _empty_stats(), "groups": {}, "has_segments": False}
                for hours in WINDOWS
            },
            "rows": [],
        }

    def report_csv(self, plan_id):
        report = self.report(plan_id)
        output = io.StringIO()
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(["行 ID", "接收人 UnionID", "话术", "发送员工 UserID", "小程序路径", "分层", "发送状态", "实际发送时间", "12 小时观察", "24 小时观察", "48 小时观察"])
        for row in report.get("rows", []):
            windows = row.get("windows")
            windows = windows if isinstance(windows, dict) else {}
            state = row.get("state") or row.get("delivery_state")
            writer.writerow([
                _safe_csv(row.get("id")),
                _safe_csv(row.get("unionid")),
                _safe_csv(row.get("text")),
                _safe_csv(row.get("sender_userid")),
                _safe_csv(row.get("path")),
                _safe_csv(row.get("segment")),
                _safe_csv(_report_label(state, DELIVERY_LABELS, "状态待核对")),
                _safe_csv(_report_time(row.get("sent_at"))),
                _safe_csv(_report_label(windows.get("12"), WINDOW_LABELS)),
                _safe_csv(_report_label(windows.get("24"), WINDOW_LABELS)),
                _safe_csv(_report_label(windows.get("48"), WINDOW_LABELS)),
            ])
        return output.getvalue().encode("utf-8")


class Handler(BaseHTTPRequestHandler):
    server_version = "ExcelBatch"

    def log_message(self, *args):
        pass  # Do not log identifiers, request bodies, or credentials.

    def do_GET(self):
        self.handle_api()

    def do_POST(self):
        self.handle_api()

    def do_PUT(self):
        self.handle_api()

    def handle_api(self):
        with self.server.service.source.session():
            self._handle_api()

    def _read_body(self, limit):
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            raise Invalid("Content-Length 无效")
        if length < 1 or length > limit:
            raise Invalid("请求体过大或为空")
        return self.rfile.read(length)

    def _handle_api(self):
        try:
            expected = "Bearer " + self.server.token
            if not hmac.compare_digest(self.headers.get("Authorization", ""), expected):
                return self.reply(401, {"error": "unauthorized"})
            parsed = urllib.parse.urlsplit(self.path)
            path = parsed.path
            service = self.server.service
            if self.command == "GET" and path == "/health":
                return self.reply(200, {"ok": True})
            if self.command == "GET" and path.startswith("/covers/"):
                key = urllib.parse.unquote(path[len("/covers/"):])
                with service.db() as db:
                    row = db.execute("SELECT body,mime FROM covers WHERE digest=?", (key,)).fetchone()
                if not row:
                    return self.reply(404, {"error": "cover_not_found"})
                return self.reply(200, row["body"], row["mime"])
            if self.command == "GET" and path.startswith("/reports/"):
                raw_id = path[len("/reports/"):]
                if raw_id.endswith(".csv"):
                    return self.reply(200, service.report_csv(int(raw_id[:-4])), "text/csv; charset=utf-8")
                return self.reply(200, service.report(int(raw_id)))
            if path == "/prepare" and self.command == "POST":
                return self.reply(200, service.prepare_import(self._read_body(MAX_BYTES)))
            if path == "/imports" and self.command == "POST":
                raw = self._read_body(MAX_BYTES)
                query = urllib.parse.parse_qs(parsed.query)
                response = service.import_file(raw, query.get("new") == ["1"], self.headers.get("Idempotency-Key", ""))
                return self.reply(200, response)
            if self.command != "POST":
                return self.reply(404, {"error": "not_found"})
            raw = self._read_body(24 * 1024 * 1024)
            if path == "/covers":
                return self.reply(200, {"cover_digest": service.cover(raw)})
            data = json.loads(raw)
            if path == "/snapshots":
                return self.reply(200, service.snapshot(data["snapshot_key"], data["rows"], data.get("segment_source"), data.get("has_segments"), data.get("version")))
            if path == "/observations":
                return self.reply(200, service.observe(int(data["plan_id"]), data["snapshot_key"], data["rows"], data.get("now")))
            if path == "/content-opens":
                return self.reply(200, service.content_open(data))
            return self.reply(404, {"error": "not_found"})
        except Conflict as error:
            body = {"error": error.code, "message": str(error)}
            return self.reply(409, body)
        except Invalid as error:
            body = {"error": "invalid_input", "message": str(error)}
            if error.row_errors:
                body["row_errors"] = error.row_errors
            return self.reply(400, body)
        except (ValueError, KeyError, TypeError):
            return self.reply(400, {"error": "invalid_input"})
        except Exception:
            return self.reply(503, {"error": "component_unavailable"})

    def reply(self, status, body, mime="application/json"):
        raw = body if isinstance(body, bytes) else json.dumps(body, ensure_ascii=False).encode()
        self.send_response(status)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(raw)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", required=True)
    parser.add_argument("--database", required=True)
    parser.add_argument("--port", type=int, default=8791)
    args = parser.parse_args()
    config = json.loads(Path(args.config).read_text())
    token = os.environ.get("EXCEL_BATCH_TOKEN", "")
    if len(token) < 32 or not config.get("appid"):
        raise SystemExit("Configure token and appid before starting")
    service = Service(args.database, config)
    server = ThreadingHTTPServer(("127.0.0.1", args.port), Handler)
    server.service, server.token = service, token
    server.serve_forever()


if __name__ == "__main__":
    main()
